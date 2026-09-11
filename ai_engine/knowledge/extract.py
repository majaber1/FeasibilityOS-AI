"""Text extraction for PDF / DOCX / XLSX knowledge uploads."""
from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Tuple


def extract_text(data: bytes, filename: str, content_type: str | None = None) -> str:
    name = (filename or "").lower()
    ctype = (content_type or "").lower()
    if name.endswith(".pdf") or "pdf" in ctype:
        return _pdf_text(data)
    if name.endswith(".docx") or "wordprocessingml" in ctype:
        return _docx_text(data)
    if name.endswith(".xlsx") or name.endswith(".xls") or "spreadsheetml" in ctype:
        return _xlsx_text(data)
    if name.endswith(".txt") or ctype.startswith("text/"):
        return data.decode("utf-8", errors="ignore")
    # Best-effort decode
    return data.decode("utf-8", errors="ignore")


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except Exception:
        return ""
    reader = PdfReader(io.BytesIO(data))
    parts: List[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts).strip()


def _docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except Exception:
        return ""
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _xlsx_text(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except Exception:
        return ""
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts: List[str] = []
    for sheet in wb.worksheets:
        parts.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals:
                parts.append(" | ".join(vals))
    return "\n".join(parts).strip()


_SAR_RE = re.compile(
    r"(?i)(?:SAR|SR|ريال)\s*([0-9]{1,3}(?:,[0-9]{3})*(?:\.[0-9]+)?|[0-9]+(?:\.[0-9]+)?)"
)
_PCT_RE = re.compile(
    r"(?i)(occupancy|absorption|irr|roi|margin|churn|utilization|load\s*factor)[^\n%]{0,40}?(\d{1,3}(?:\.\d+)?)\s*%"
)
_MW_RE = re.compile(r"(?i)(\d+(?:\.\d+)?)\s*MW")
_PROJECT_HINTS = [
    ("data_center", re.compile(r"(?i)data\s*center|colocation|PUE|rack\s*power|تيرابايت|مركز\s*بيانات")),
    ("real_estate", re.compile(r"(?i)residential|compound|villa|apartment|absorption|عقار|مجمع\s*سكني")),
    ("saas_digital", re.compile(r"(?i)\bSaaS\b|ARR|MRR|churn|CAC|LTV|اشتراك")),
    ("services", re.compile(r"(?i)cyber|MSSP|consulting|managed\s*service|SOC|أمن\s*سيبراني")),
    ("industrial", re.compile(r"(?i)factory|manufactur|plant|صناع")),
    ("retail", re.compile(r"(?i)retail|store|mall|تجزئة")),
]


def extract_structured_metadata(text: str, filename: str = "") -> Dict[str, Any]:
    blob = f"{filename}\n{text or ''}"
    project_type = None
    for key, rx in _PROJECT_HINTS:
        if rx.search(blob):
            project_type = key
            break

    amounts = [float(m.group(1).replace(",", "")) for m in _SAR_RE.finditer(blob)]
    capex = amounts[0] if amounts else None
    opex = amounts[1] if len(amounts) > 1 else None

    assumptions: Dict[str, Any] = {}
    for m in _PCT_RE.finditer(blob):
        assumptions[m.group(1).lower().replace(" ", "_")] = f"{m.group(2)}%"

    mw = _MW_RE.search(blob)
    if mw:
        assumptions["it_load_mw"] = mw.group(1)

    year = None
    ym = re.search(r"\b(20[2-3][0-9])\b", blob)
    if ym:
        year = int(ym.group(1))

    sector = None
    if project_type == "data_center":
        sector = "technology_infrastructure"
    elif project_type == "real_estate":
        sector = "real_estate"
    elif project_type == "saas_digital":
        sector = "software"
    elif project_type == "services":
        sector = "professional_services"

    outcome = {}
    if re.search(r"(?i)\b(GO|feasible|viable|موصى|جدوى\s*إيجاب)", blob):
        outcome["signal"] = "positive"
    elif re.search(r"(?i)\b(NO[\s-]?GO|not feasible|غير\s*مجدي)", blob):
        outcome["signal"] = "negative"

    confidence = 0.35
    if project_type:
        confidence += 0.2
    if assumptions:
        confidence += 0.15
    if amounts:
        confidence += 0.1
    confidence = min(confidence, 0.85)

    return {
        "project_type": project_type,
        "sector": sector,
        "year": year,
        "country": "SA",
        "capex": {"amount": capex, "currency": "SAR"} if capex is not None else None,
        "opex": {"amount": opex, "currency": "SAR"} if opex is not None else None,
        "revenue_model": _guess_revenue_model(project_type, blob),
        "assumptions": assumptions,
        "outcome": outcome or None,
        "confidence": confidence,
    }


def _guess_revenue_model(project_type: str | None, blob: str) -> Dict[str, Any] | None:
    if project_type == "data_center":
        return {"type": "colocation_lease", "unit": "per_kW_month"}
    if project_type == "real_estate":
        return {"type": "unit_sales_and_lease"}
    if project_type == "saas_digital":
        return {"type": "subscription"}
    if project_type == "services":
        return {"type": "retainer_and_projects"}
    if re.search(r"(?i)take[\s-]?rate|commission", blob):
        return {"type": "marketplace_take_rate"}
    return None
